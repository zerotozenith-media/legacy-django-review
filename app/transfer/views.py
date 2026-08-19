from datetime import datetime
from django.utils import timezone
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Sum
from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import ListView
from authentication.middleware import Role
from organization.models import BankAccount, LocalChurch, Region, BankAccount
from organization.utils import get_user_local_church
from .models import Transfer
from django.core.paginator import Paginator
from django.template import loader
from django.contrib import messages
from django.contrib.auth.decorators import login_required
import json
from django.http import JsonResponse
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
import re
from django.db import transaction
from django.http import JsonResponse, HttpResponse
from openpyxl import load_workbook, Workbook
from contributor.helper import validate_date, check_amount
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
import re
from django.db import transaction
from django.http import JsonResponse, HttpResponse
from openpyxl import load_workbook, Workbook
from contributor.helper import validate_date, check_amount
from django.views import View


def current_month_transfer(user=None):
    local_church = []
    role = user.role
    current_year = datetime.now().year
    current_month = datetime.now().month
    total_transfer_current_month = Transfer.objects.filter(
        date__month=current_month, date__year=current_year
    )

    if role == Role.region_overseer.value:
        regions = user.managed_region.all()
        for region in regions:
            local_church += region.local_churches.all()
    elif role == Role.region_overseer.value:
        local_church = user.manager_local_church.all()
    elif role == Role.national_overseer.value or user.is_superuser:
        local_church = LocalChurch.objects.all()

    if role == Role.local_church_pastor.value and local_church:
        total_transfer_current_month = total_transfer_current_month.filter(local_church__in=local_church)
    total_transfer_current_month = total_transfer_current_month.aggregate(total=Sum('amount'))['total'] or 0.0
    return total_transfer_current_month


def current_year_transfer(user=None):
    local_church = []
    role = user.role
    current_year = datetime.now().year
    total_transfer_current_year = Transfer.objects.filter(
        date__year=current_year
    )
    if role == Role.region_overseer.value:
        regions = user.managed_region.all()
        for region in regions:
            local_church += region.local_churches.all()
    elif role == Role.region_overseer.value:
        local_church = user.manager_local_church.all()
    elif role == Role.national_overseer.value or user.is_superuser:
        local_church = LocalChurch.objects.all()

    if role == Role.local_church_pastor.value and local_church:
        total_transfer_current_year = total_transfer_current_year.filter(local_church__in=local_church)
    total_transfer_current_year = total_transfer_current_year.aggregate(total=Sum('amount'))['total'] or 0.0
    return total_transfer_current_year


def search_transfer(request):
    if request.method == "POST":
        search_str = json.loads(request.body).get("searchText")
        transfer = (
                Transfer.objects.filter(amount__istartswith=search_str, owner=request.user)
                | Transfer.objects.filter(date__istartswith=search_str, owner=request.user)
                | Transfer.objects.filter(description__icontains=search_str, owner=request.user)
        )
        data = transfer.values()
        return JsonResponse(list(data), safe=False)


@login_required(login_url="/authentication/login")
def index(request):
    start_date = request.GET.get('start_date', None)
    end_date = request.GET.get('end_date', None)
    # paginator = Paginator(transfer, 10)
    page_number = request.GET.get("page")
    # page_obj = Paginator.get_page(paginator, page_number)
    # currency = UserPreference.objects.get(user=request.user).currency
    user = request.user
    if user.is_staff or user.role == 3:
        all_regions = Region.objects.all().order_by('name')
    else:
        all_regions = Region.objects.filter(region_overseer=user).order_by('name')
    transfer = Transfer.objects.filter(region__in=all_regions)
    if start_date and end_date:
        if isinstance(start_date, str):
            start_date = datetime.strptime(start_date, "%d-%m-%Y").date()
        if isinstance(end_date, str):
            end_date = datetime.strptime(end_date, "%d-%m-%Y").date()
        transfer = transfer.filter(date__range=[start_date, end_date], deleted=False)
    else:
        current_date = timezone.now().date()
        transfer = transfer.filter(date__range=[current_date, current_date], deleted=False)
    transfer_this_month = current_month_transfer(user=request.user)
    transfer_this_year = current_year_transfer(user=request.user)

    context = {
        "transfer": transfer,
        "page_obj": transfer,
        "current_month_transfer": transfer_this_month,
        "current_year_transfer": transfer_this_year,
        # 'currency': currency
    }
    return render(request, "transfer/index.html", context)


@login_required(login_url="/authentication/login")
def add_transfer(request):
    accounts = BankAccount.objects.all()
    user = request.user
    if user.is_staff or user.role == 3:
        all_regions = Region.objects.all().order_by('name')
    else:
        all_regions = Region.objects.filter(region_overseer=user).order_by('name')
    context = {
        "all_regions": all_regions,
        "accounts": accounts,
        "values": request.POST,
        "current_user": int(request.user.id)
    }
    if request.method == "GET":
        return render(request, "transfer/add_transfer.html", context)

    if request.method == "POST":
        print(request.POST)
        region_id = request.POST["region_id"]
        account = request.POST["account"]
        amount = request.POST["amount"]
        transfer_type = request.POST["type"]
        description = request.POST["description"]
        date = request.POST["transfer_date"]

        if not region_id:
            messages.error(request, "Transfer region required")
            return render(request, "transfer/add_transfer.html", context)

        if not amount:
            messages.error(request, "Amount is required")
            return render(request, "transfer/add_transfer.html", context)

        if not description:
            messages.error(request, "Description is required")
            return render(request, "transfer/add_transfer.html", context)

        if not date:
            messages.error(request, "Date is required")
            return render(request, "transfer/add_transfer.html", context)


        Transfer.objects.create(
            region_id=region_id,
            account_id=account,
            amount=amount,
            date=date,
            type=transfer_type,
            description=description,
        )
        messages.success(request, "transfer saved successfully")

        return redirect("transfer")


@login_required(login_url="/authentication/login")
def transfer_edit(request, id):
    transfer = get_object_or_404(Transfer, pk=id)
    all_regions = Region.objects.all()  
    accounts = BankAccount.objects.all()  
    
    if request.method == "POST":
        transfer.type = request.POST.get("type")
        transfer.region_id = request.POST.get("region_id")
        transfer.account_id = request.POST.get("account")
        transfer.date = request.POST.get("transfer_date")
        transfer.amount = request.POST.get("amount")
        transfer.description = request.POST.get("description")

        # Validation checks
        if not transfer.amount:
            messages.error(request, "Amount is required")
        elif not transfer.description:
            messages.error(request, "Description is required")
        else:
            transfer.save()
            messages.success(request, "Transfer record updated successfully!")
            return redirect("transfer")
    
    context = {
        "transfer": transfer,
        "all_regions": all_regions,
        "accounts": accounts,
    }
    return render(request, "transfer/edit_transfer.html", context)

def delete_transfer(request, id):
    transfer = Transfer.objects.get(pk=id)
    transfer.delete()
    messages.success(request, "Record was removed")
    return redirect("transfer")


class TransferUploaderHelper:
    @staticmethod
    def validate_headers(worksheet):
        headers = [cell.value.strip() if cell.value else None for cell in worksheet[1]]
        return headers == ['Account', 'Amount', 'Description', 'Date', 'Type']

    @staticmethod
    def process_and_save(worksheet, region_id):
        with transaction.atomic():
            for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                account_name, amount, description, date, type = row

                if not account_name and not amount and not description and not date:
                    continue

                region = Region.objects.get(id=region_id)
                print(account_name)
                account = BankAccount.objects.get(account_name=account_name)

                Transfer.objects.create(
                    region=region,
                    account=account,
                    amount=amount,
                    description=description,
                    date=date,
                    type=type,
                )


class TransferUploadView(View):
    @staticmethod
    def validate_transfer_uploader(file, region_id):
        transfer_upload_helper = TransferUploaderHelper()
        try:
            workbook = load_workbook(file, read_only=True)
            worksheet = workbook.active

            if not transfer_upload_helper.validate_headers(worksheet):
                return {'error': 'Invalid headers'}

            for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                account_name, amount, description, date, type = row

                if not account_name and not amount and not description and not date and not type:
                    continue

                if date and not validate_date(date):
                    return {'error': f'Invalid date format: "{date}" in Row: {row_number}'}

                if amount and not check_amount(amount):
                    return {'error': f'Invalid amount format: "{amount}" in Row: {row_number}'}

                if not amount or not date or not account_name:
                    data = 'amount' if not amount else 'date' if not date else 'account_name'
                    return {'error': f'Required field missing in Row {row_number}. Please provide {data}.'}

            transfer_upload_helper.process_and_save(worksheet, region_id)

        except Exception as e:
            return {'error': f'Error processing the file: {str(e)}'}

    def get(self, request):
        all_regions = []
        user = request.user
        if user.is_staff or user.role == 3:
            all_regions = Region.objects.all().order_by('name')
        else:
            all_regions = Region.objects.filter(region_overseer=user).order_by('name')
        return render(request, 'transfer/transfer_uploader.html', {'all_regions': all_regions})

    def post(self, request):
        errors = self.validate_transfer_uploader(request.FILES['file'], request.POST.get('region_id'))
        if errors:
            all_regions = []
            user = request.user
            if user.is_staff or user.role == 3:
                all_regions = Region.objects.all().order_by('name')
            else:
                all_regions = Region.objects.filter(region_overseer=user).order_by('name')
            return render(request, 'transfer/transfer_uploader.html', {'errors': errors, 'all_regions': all_regions})
        messages.success(request, 'Transfer records uploaded and processed successfully!')
        return redirect('transfer')


def download_transfer_uploader_sample(request):
    workbook = Workbook()
    worksheet = workbook.active

    headers = ['Account', 'Amount', 'Description', 'Date', 'Type']
    worksheet.append(headers)

    header_row = worksheet[1]
    for cell in header_row:
        cell.font = Font(bold=True)

    data = [
        ['Saskatoon', '1500.75', 'Transfer Description', '2024-08-01','Received'],
        ['Regina', '250.00', 'Another Transfer', '2024-07-15','Sent'],
        ['Ottawa', '350.00', 'Another Transfer', '2024-07-18','Sent'],
    ]
    for sample_data in data:
        worksheet.append(sample_data)

    for col_num, column_letter in enumerate(get_column_letter(len(headers)), start=1):
        worksheet.column_dimensions[column_letter].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=sample_transfer_uploader_file.xlsx'
    workbook.save(response)

    return response
